"""Check report figures, stored numerical values and editable Word content."""

import unittest

import pandas as pd
from docx import Document

from monthly_backtest import ROOT, OUTPUT


class ReportTests(unittest.TestCase):
    def test_report_numbers_equal_csv_values(self):
        markdown = (ROOT / "report/Report.md").read_text(encoding="utf-8")
        doc = Document(ROOT / "report/Report.docx")
        cells = [
            cell.text
            for table in doc.tables
            for row in table.rows
            for cell in row.cells
        ]
        word_text = "\n".join([p.text for p in doc.paragraphs] + cells)
        metrics = pd.read_csv(OUTPUT / "period_comparison_metrics.csv")
        for row in metrics.itertuples():
            pair = f"{row.MAE_kW:.6f} / {row.RMSE_kW:.6f}"
            self.assertIn(pair, markdown)
            self.assertIn(pair, cells)
        for line in markdown.splitlines():
            line = line.strip()
            if not line or line.startswith(("![", "|")):
                continue
            paragraph = line.lstrip("# ") if line.startswith("#") else line
            self.assertIn(paragraph, word_text)
        self.assertEqual(doc.styles["Normal"].font.name, "Times New Roman")
        self.assertEqual(doc.styles["Normal"].font.size.pt, 11)
        self.assertEqual(doc.styles["Normal"].paragraph_format.line_spacing, 1.5)
        self.assertEqual(doc.core_properties.comments, "")
        self.assertEqual(doc.core_properties.author, "")
        self.assertEqual(doc.core_properties.last_modified_by, "")
        self.assertGreater(doc.core_properties.created.year, 2013)
        self.assertEqual(doc.core_properties.created, doc.core_properties.modified)
        self.assertTrue(doc.styles["Heading 1"].paragraph_format.keep_with_next)
        backtest_paragraph = next(
            p for p in doc.paragraphs if p.text.startswith("Experiment A fits ")
        )
        self.assertTrue(backtest_paragraph.paragraph_format.keep_together)
        self.assertIn("SSE (kW²)", markdown)
        self.assertIn("SSE (kW²)", cells)
        self.assertNotIn("SSE (kW2)", markdown)
        errors = pd.read_csv(ROOT / "results/july_analysis/error_distribution.csv")
        for value in errors.sum_squared_error_kW2:
            self.assertIn(f"{value:.3f}", cells)

    def test_three_figures_and_formats_exist(self):
        for name in [
            "figure1_consumption_pattern.png",
            "figure2_july_forecast.png",
            "figure3_monthly_performance.png",
            "figure2_forecast_comparison.png",
        ]:
            self.assertTrue((ROOT / "figures" / name).exists())
        self.assertEqual(len(Document(ROOT / "report/Report.docx").inline_shapes), 3)
        for extension in ["md", "pdf", "docx"]:
            self.assertTrue((ROOT / f"report/Report.{extension}").is_file())


if __name__ == "__main__":
    unittest.main()
