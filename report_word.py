from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent


def generate_word():
    """Generate the editable Word report from the canonical Markdown report."""
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.65)
    sec.left_margin = sec.right_margin = Inches(0.75)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Caption"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name in ["Normal", "Caption"] else 14)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.space_before = Pt(8 if name.startswith("Heading") else 0)
        style.font.bold = name in ["Title", "Heading 1", "Heading 2"]
        style.font.italic = False
        if name.startswith("Heading"):
            style.paragraph_format.keep_with_next = True
            style.paragraph_format.keep_together = True
        style.element.get_or_add_rPr()
        fonts = style.element.rPr.rFonts
        if fonts is None:
            fonts = OxmlElement("w:rFonts")
            style.element.rPr.append(fonts)
        for key in ["ascii", "hAnsi", "eastAsia", "cs"]:
            fonts.set(qn("w:" + key), "Times New Roman")
            fonts.attrib.pop(qn("w:" + key + "Theme"), None)
    for border in doc.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    lines = (ROOT / "report/Report.md").read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(c == "---" for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=len(rows[0]))
            table.style = "Table Grid"
            for ri, row in enumerate(rows):
                table_row = table.add_row()
                table_row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                cells = table_row.cells
                for c, t in zip(cells, row):
                    c.text = t
                    for p in c.paragraphs:
                        p.paragraph_format.line_spacing = 1.15
                        p.paragraph_format.keep_with_next = ri < len(rows) - 1
                        for run in p.runs:
                            run.font.name = "Times New Roman"
                            run.font.size = Pt(10)
                            run.bold = ri == 0
            continue
        if line.startswith("!["):
            image = line.split("](")[1].rstrip(")")
            doc.add_picture(str((ROOT / "report" / image).resolve()), width=Inches(6.6))
            doc.paragraphs[-1].paragraph_format.keep_with_next = True
        elif line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], style="Heading 2")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        else:
            paragraph = doc.add_paragraph(
                line,
                style=(
                    "Caption"
                    if line.startswith(("Figure 1.", "Figure 2.", "Figure 3."))
                    else None
                ),
            )
            if line.startswith("Table "):
                paragraph.paragraph_format.keep_with_next = True
            if line.startswith("Experiment A fits "):
                paragraph.paragraph_format.keep_together = True
            if line.startswith(("[1]", "[2]", "[3]", "[4]")):
                paragraph.paragraph_format.left_indent = Inches(0.25)
                paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        i += 1
    doc.core_properties.title = next(
        line[2:] for line in lines if line.startswith("# ")
    )
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    generated_at = datetime.now(timezone.utc).replace(microsecond=0)
    doc.core_properties.created = generated_at
    doc.core_properties.modified = generated_at
    doc.save(ROOT / "report/Report.docx")


if __name__ == "__main__":
    generate_word()
